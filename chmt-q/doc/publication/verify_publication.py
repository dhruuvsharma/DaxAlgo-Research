"""Verify CHMT-Q manuscript claims, artifacts, DOCX structure, and PDF rendering."""

from __future__ import annotations

import hashlib
import json
import math
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import fitz
from docx import Document


HERE = Path(__file__).resolve().parent
MANUSCRIPT = HERE / "manuscript.md"
DOCX = HERE / "chmtq-negative-results-manuscript.docx"
PDF = HERE / "chmtq-negative-results-manuscript.pdf"
ASSETS = HERE / "assets"

STUDY = Path(r"C:\DaxAlgoData\okx\BTC-USDT\2026-06-chmtq-research-v1")
BACKTEST = STUDY / "backtest-fold1-step2010-v1"
DIAGNOSTICS = STUDY / "failure-analysis-fold1-step2010-v4" / "diagnostics.json"
SUPPLEMENT = HERE.parent / "evidence/failure_revision_supplement.json"
DQN_PLAN = HERE.parent / "evidence/transformer_dqn_architecture_plan.json"

EXPECTED_HASHES = {
    BACKTEST / "results.json": "083da170146585cf1a16dde56e2da02697194640ed7b053d3298d8e3775bee6a",
    BACKTEST / "study-protocol.json": "885fcd0fee23115a4a5480ec94d95f436e251ec0caf796f9686867ed7e72f5d0",
    BACKTEST / "predictions.jsonl.gz": "27d9ffbe41b18868b8a45bf9c48c936908de7394dd6ecfec7062cc531cc7e1f0",
    DIAGNOSTICS: "ece5b18f8a8455f7edac8491b6d23362fa14f618129b1193b1febf14bc89505e",
    DQN_PLAN: "d5e81f203619c55754c940df7e92ec5995e4db41b3a7d6f3db4137318f7a502c",
}


def main() -> int:
    checks: dict[str, object] = {}
    verify_source(checks)
    verify_evidence(checks)
    verify_docx(checks)
    verify_pdf(checks)
    checks["status"] = "pass"
    print(json.dumps(checks, indent=2, sort_keys=True))
    return 0


def require(condition: bool, message: str) -> None:
    if not condition:
        raise AssertionError(message)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def verify_source(checks: dict[str, object]) -> None:
    text = MANUSCRIPT.read_text(encoding="utf-8")
    expected = [
        "2,089,554-parameter",
        "0.251391",
        "0.257193",
        "2.2558%",
        "53.89",
        "reinforcement-learning design is a preregistered future-work architecture",
        "process-level holdout policy",
        "No sealed values were printed, retained, analyzed, or used",
    ]
    for phrase in expected:
        require(phrase in text, f"Required claim missing: {phrase}")
    forbidden = [
        "Jun 25-30 holdout unaccessed",
        "production-ready model",
        "proven profitable",
        "statistically significant improvement",
        "DQN was trained",
        "DQN results",
        "beats strong baselines",
    ]
    for phrase in forbidden:
        require(phrase.casefold() not in text.casefold(), f"Unsafe claim present: {phrase}")

    marker_figures = re.findall(r"^\{FIGURE:", text, flags=re.MULTILINE)
    table_separators = re.findall(r"^\|(?:\s*:?-{3,}:?\s*\|)+$", text, flags=re.MULTILINE)
    references = re.findall(r"^(\d+)\.\s", text.split("## References", 1)[1].split("## Appendix", 1)[0], flags=re.MULTILINE)
    cited: set[int] = set()
    front_matter = text.split("## References", 1)[0]
    for match in re.finditer(r"\[([\d,–\-\s]+)\]", front_matter):
        for token in match.group(1).replace(" ", "").split(","):
            if "–" in token or "-" in token:
                parts = re.split(r"[–-]", token)
                if len(parts) == 2 and all(part.isdigit() for part in parts):
                    cited.update(range(int(parts[0]), int(parts[1]) + 1))
            elif token.isdigit():
                cited.add(int(token))
    require(len(marker_figures) == 8, "Expected eight figure markers")
    require(len(table_separators) == 10, "Expected ten Markdown tables")
    require(references == [str(index) for index in range(1, 31)], "References must be sequential 1–30")
    require(set(range(1, 31)).issubset(cited), f"Uncited references: {sorted(set(range(1, 31)) - cited)}")
    checks["source"] = {
        "words": len(re.findall(r"\b\w+[\w–-]*\b", text)),
        "figures": len(marker_figures),
        "tables": len(table_separators),
        "references": len(references),
    }


def verify_evidence(checks: dict[str, object]) -> None:
    for path, expected in EXPECTED_HASHES.items():
        require(path.is_file(), f"Missing evidence artifact: {path}")
        actual = sha256(path)
        require(actual == expected, f"Hash mismatch: {path}")

    diagnostics = json.loads(DIAGNOSTICS.read_text(encoding="utf-8"))
    score = diagnostics["scoring_diagnostics"]["untouched_development_forward"]["weighted_normalized_pinball"]
    h5 = diagnostics["partition_metrics"]["untouched_development_forward"]["cells"][4][0]
    economics = diagnostics["economic_diagnostics"]
    contract = diagnostics["model_contract"]
    training = diagnostics["training_diagnostics"]

    exact = [
        (contract["parameter_count"], 2_089_554),
        (contract["train_rows"], 25_758),
        (score["model"], 0.2513909348275368),
        (score["empirical_baseline"], 0.25719276632184623),
        (score["model_improvement_percent"], 2.2558299664809045),
        (h5["median_mae"], 532.8411673773093),
        (h5["pearson"], 0.0774449555902145),
        (h5["spearman"], 0.07533437940317818),
        (h5["prediction_to_actual_std_ratio"], 0.15129684258798334),
        (h5["interval_crosses_zero_fraction"], 1.0),
        (economics["median_long_flat"]["trade_count"], 484),
        (economics["median_long_flat"]["base_cost_compounded_return"], -0.680682183408049),
        (economics["median_long_flat"]["base_cost_to_break_even_multiple"], 53.89286674073603),
        (training["best_epoch"], 10),
        (training["completed_optimizer_steps"], 3618),
    ]
    for actual, expected in exact:
        require(math.isclose(float(actual), float(expected), rel_tol=0, abs_tol=1e-12), f"Evidence mismatch: {actual} != {expected}")

    require(diagnostics["outer_holdout_used_in_failure_metrics"] is False, "Failure metrics used outer holdout")
    supplement_text = SUPPLEMENT.read_text(encoding="utf-8")
    require("inadvertently_read_not_used" in supplement_text, "Holdout integrity status changed")
    plan = json.loads(DQN_PLAN.read_text(encoding="utf-8"))
    require(plan["status"] == "proposed_not_implemented_not_validated", "DQN status is no longer proposal-only")
    checks["evidence"] = {"hashes_verified": len(EXPECTED_HASHES), "metrics_verified": len(exact)}


def verify_docx(checks: dict[str, object]) -> None:
    require(DOCX.is_file() and DOCX.stat().st_size > 1_000_000, "DOCX missing or unexpectedly small")
    document = Document(DOCX)
    heading_counts: dict[str, int] = {}
    for paragraph in document.paragraphs:
        if paragraph.style and paragraph.style.name.startswith("Heading"):
            heading_counts[paragraph.style.name] = heading_counts.get(paragraph.style.name, 0) + 1
    require(heading_counts.get("Heading 1", 0) >= 15, "Heading hierarchy incomplete")
    require(len(document.inline_shapes) == 8, "DOCX must contain eight inline figures")
    require(len(document.tables) == 10, "DOCX must contain ten data tables")

    namespaces = {
        "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
        "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
        "cp": "http://schemas.openxmlformats.org/package/2006/metadata/core-properties",
        "dc": "http://purl.org/dc/elements/1.1/",
    }
    with zipfile.ZipFile(DOCX) as archive:
        xml = ET.fromstring(archive.read("word/document.xml"))
        descriptions = [node.attrib.get("descr", "").strip() for node in xml.findall(".//wp:docPr", namespaces)]
        require(len(descriptions) == 8 and all(descriptions), "Every figure needs non-empty alt text")
        require(not xml.findall(".//w:ins", namespaces), "Tracked insertions remain")
        require(not xml.findall(".//w:del", namespaces), "Tracked deletions remain")
        for table in xml.findall(".//w:tbl", namespaces):
            first_row = table.find("w:tr", namespaces)
            require(first_row is not None and first_row.find("w:trPr/w:tblHeader", namespaces) is not None, "Data table lacks a header row")
        core = ET.fromstring(archive.read("docProps/core.xml"))
        creators = [node.text or "" for node in core.findall("dc:creator", namespaces)]
        require(all(not value.strip() or value.strip() == "Anonymous" for value in creators), "Identifying author metadata remains")

    checks["docx"] = {
        "bytes": DOCX.stat().st_size,
        "sha256": sha256(DOCX),
        "headings": heading_counts,
        "images_with_alt": 8,
        "tables_with_headers": 10,
    }


def verify_pdf(checks: dict[str, object]) -> None:
    require(PDF.is_file() and PDF.stat().st_size > 500_000, "PDF missing or unexpectedly small")
    document = fitz.open(PDF)
    require(len(document) == 21, f"Unexpected PDF page count: {len(document)}")
    require(not (document.metadata.get("author") or "").strip(), "PDF author metadata was not anonymized")
    all_text = "\n".join(page.get_text("text") for page in document)
    for phrase in [
        "When Better Proper Scores Do Not Imply Tradable Edge",
        "2,089,554",
        "2.2558%",
        "53.89",
        "References",
        "Appendix C. Submission Checklist",
    ]:
        require(phrase in all_text, f"PDF text missing: {phrase}")
    require("�" not in all_text, "PDF contains replacement glyphs")
    require("□" not in all_text, "PDF contains missing-glyph boxes")

    blank_pages = []
    outside = []
    links = 0
    for page_number, page in enumerate(document, start=1):
        if len(page.get_text("text").strip()) < 80:
            blank_pages.append(page_number)
        page_rect = page.rect
        links += len(page.get_links())
        for item in page.get_text("dict").get("blocks", []):
            x0, y0, x1, y1 = item["bbox"]
            if x0 < -1 or y0 < -1 or x1 > page_rect.width + 1 or y1 > page_rect.height + 1:
                outside.append((page_number, item["bbox"]))
    require(not blank_pages, f"Nearly blank PDF pages: {blank_pages}")
    require(not outside, f"PDF content outside page bounds: {outside[:3]}")
    require(links >= 30, f"Expected at least 30 reference links, found {links}")
    pages = len(document)
    document.close()
    checks["pdf"] = {
        "bytes": PDF.stat().st_size,
        "sha256": sha256(PDF),
        "pages": pages,
        "links": links,
        "author_metadata": "blank",
        "glyph_check": "pass",
        "bounds_check": "pass",
    }


if __name__ == "__main__":
    raise SystemExit(main())
