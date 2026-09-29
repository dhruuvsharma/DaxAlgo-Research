"""Build the venue-neutral CHMT-Q publication manuscript and its figures."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path

import matplotlib

matplotlib.use("Agg")

import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "manuscript.md"
ASSETS = ROOT / "assets"
OUTPUT = ROOT / "chmtq-negative-results-manuscript.docx"

INK = "172033"
NAVY = "102A43"
TEAL = "0F766E"
BLUE = "2F5D8C"
RUST = "B44C3B"
AMBER = "C88918"
MUTED = "5F6B78"
GRID = "CBD5E1"
PALE = "EEF3F7"
PALE_TEAL = "E7F4F1"
PALE_RED = "F9ECE9"
WHITE = "FFFFFF"

TABLE_CAPTIONS = [
    "Claim levels, registered questions, and defensible conclusions.",
    "Implemented CHMT-Q tensor and architecture contract.",
    "Training, selection, and artifact summary.",
    "Development-forward forecasting endpoints.",
    "Economic outcomes under the frozen execution proxy.",
    "Ranked failure factors and evidence status.",
    "Eligibility gates for the proposed Transformer–Double-DQN study.",
    "Immutable artifact and experiment-identity ledger.",
    "Tensor and parameter ledger.",
    "Registration status of reported analyses.",
]


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=SOURCE)
    parser.add_argument("--output", type=Path, default=OUTPUT)
    args = parser.parse_args()

    source = args.source.resolve(strict=True)
    output = args.output.resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    ASSETS.mkdir(parents=True, exist_ok=True)

    build_figures(ASSETS)
    doc = build_document(source)
    if output.exists():
        output.unlink()
    doc.save(output)
    print(
        json.dumps(
            {
                "output": str(output),
                "bytes": output.stat().st_size,
                "sha256": sha256(output),
                "figures": len(list(ASSETS.glob("*.png"))),
            },
            sort_keys=True,
        )
    )
    return 0


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def build_figures(output_dir: Path) -> None:
    plt.rcParams.update(
        {
            "font.family": "DejaVu Sans",
            "font.size": 10.5,
            "axes.titlesize": 15,
            "axes.labelsize": 10,
            "axes.edgecolor": "#9AA6B2",
            "axes.linewidth": 0.8,
            "xtick.color": "#465362",
            "ytick.color": "#465362",
            "text.color": "#172033",
            "figure.facecolor": "white",
            "axes.facecolor": "white",
            "savefig.facecolor": "white",
        }
    )
    figure_evidence_map(output_dir / "evidence_map.png")
    figure_timeline(output_dir / "study_timeline.png")
    figure_architecture(output_dir / "architecture.png")
    figure_score_heatmap(output_dir / "score_heatmap.png")
    figure_h5_diagnostics(output_dir / "h5_diagnostics.png")
    figure_cost_feasibility(output_dir / "cost_feasibility.png")
    figure_failure_map(output_dir / "failure_map.png")
    figure_next_study(output_dir / "next_study.png")


def save_figure(fig: plt.Figure, path: Path) -> None:
    fig.savefig(path, dpi=240, bbox_inches="tight", pad_inches=0.18)
    plt.close(fig)


def rect(
    ax: plt.Axes,
    xy: tuple[float, float],
    width: float,
    height: float,
    text: str,
    *,
    face: str = "#EEF3F7",
    edge: str = "#9AA6B2",
    text_color: str = "#172033",
    fontsize: float = 10,
    weight: str = "normal",
    linewidth: float = 1.2,
    hatch: str | None = None,
) -> None:
    patch = plt.Rectangle(
        xy,
        width,
        height,
        facecolor=face,
        edgecolor=edge,
        linewidth=linewidth,
        hatch=hatch,
        zorder=2,
    )
    ax.add_patch(patch)
    ax.text(
        xy[0] + width / 2,
        xy[1] + height / 2,
        text,
        ha="center",
        va="center",
        fontsize=fontsize,
        color=text_color,
        weight=weight,
        linespacing=1.25,
        zorder=3,
    )


def arrow(
    ax: plt.Axes,
    start: tuple[float, float],
    end: tuple[float, float],
    *,
    color: str = "#607080",
    style: str = "-",
) -> None:
    ax.annotate(
        "",
        xy=end,
        xytext=start,
        arrowprops={
            "arrowstyle": "-|>",
            "color": color,
            "linewidth": 1.4,
            "linestyle": style,
            "shrinkA": 2,
            "shrinkB": 2,
        },
        zorder=1,
    )


def figure_evidence_map(path: Path) -> None:
    fig, ax = plt.subplots(figsize=(12, 5.6))
    ax.set_xlim(0, 12)
    ax.set_ylim(0, 6)
    ax.axis("off")
    ax.text(0, 5.75, "What the evidence supports—and what it does not", fontsize=17, weight="bold")
    ax.text(0, 5.35, "Observed development results are separated from retired and proposed evidence.", color="#5F6B78")

    rect(ax, (0.1, 3.05), 2.2, 1.35, "OKX June 2026\nlocal research corpus\n30 UTC sessions", face="#E8EFF7", edge="#2F5D8C", weight="bold")
    rect(ax, (2.75, 3.05), 2.2, 1.35, "Implemented CHMT-Q\n2,089,554 parameters\nfootprint quantiles", face="#E7F4F1", edge="#0F766E", weight="bold")
    rect(ax, (5.4, 3.05), 2.2, 1.35, "Development-forward\n4 sessions / 5,724 rows\nfrozen checkpoint", face="#EEF3F7", edge="#607080", weight="bold")
    rect(ax, (8.05, 3.05), 3.7, 1.35, "Bounded findings\n+2.26% proper-score improvement\nno cost-feasible policy", face="#F9ECE9", edge="#B44C3B", weight="bold")
    arrow(ax, (2.3, 3.72), (2.75, 3.72))
    arrow(ax, (4.95, 3.72), (5.4, 3.72))
    arrow(ax, (7.6, 3.72), (8.05, 3.72))

    rect(ax, (2.75, 0.65), 3.25, 1.25, "Retired Jun 25–30 period\nprocess-accessed target fields\nNO CONFIRMATORY CLAIM", face="#F7F7F7", edge="#B44C3B", text_color="#7A2E23", weight="bold", hatch="///")
    rect(ax, (7.05, 0.65), 3.7, 1.25, "Proposed Transformer–Double-DQN\nunimplemented / untrained\nNO PERFORMANCE CLAIM", face="#F7F7F7", edge="#C88918", text_color="#72530C", weight="bold", hatch="\\\\")
    arrow(ax, (1.2, 3.05), (3.2, 1.9), color="#B44C3B", style="--")
    arrow(ax, (9.9, 3.05), (8.9, 1.9), color="#C88918", style="--")
    ax.text(0.1, 0.1, "Claim rule: forecast score ≠ conditional signal ≠ executable net return.", fontsize=11.5, weight="bold", color="#102A43")
    save_figure(fig, path)


def figure_timeline(path: Path) -> None:
    fig, ax = plt.subplots(figsize=(12, 4.25))
    ax.set_xlim(0.5, 30.5)
    ax.set_ylim(0, 2.2)
    ax.axis("off")
    ax.text(0.5, 2.02, "Chronological partition and current evidence status", fontsize=17, weight="bold")
    segments = [
        (1, 18, "Train replay\n18 sessions / 25,758", "#2F5D8C", "white"),
        (19, 20, "Selection\n2 / 2,862", "#C88918", "white"),
        (21, 24, "Development-forward\n4 / 5,724", "#0F766E", "white"),
        (25, 30, "Retired outer period\n6 / 8,586", "#F9ECE9", "#7A2E23"),
    ]
    for start, end, label, color, text_color in segments:
        width = end - start + 1
        patch = plt.Rectangle((start - 0.5, 0.8), width, 0.72, facecolor=color, edgecolor="white", linewidth=1.5)
        if start == 25:
            patch.set_hatch("///")
            patch.set_edgecolor("#B44C3B")
        ax.add_patch(patch)
        ax.text((start + end) / 2, 1.16, label, ha="center", va="center", color=text_color, fontsize=10.2, weight="bold")
    for day in [1, 5, 10, 15, 20, 25, 30]:
        ax.text(day, 0.58, f"Jun {day}", ha="center", va="top", color="#5F6B78", fontsize=9)
    ax.text(0.5, 0.12, "Jun 25–30: later projection read target-bearing fields; no values were retained or analyzed, but the holdout is retired.", fontsize=10.5, color="#7A2E23", weight="bold")
    save_figure(fig, path)


def figure_architecture(path: Path) -> None:
    fig, ax = plt.subplots(figsize=(12, 7.0))
    ax.set_xlim(0, 12)
    ax.set_ylim(0, 7.2)
    ax.axis("off")
    ax.text(0, 6.95, "CHMT-Q v1: implemented multimodal forecast architecture", fontsize=17, weight="bold")
    ax.text(0, 6.55, "Causal denotes temporal masking and feature availability—not causal inference.", color="#5F6B78")

    input_y = [5.15, 3.95, 2.75, 1.55]
    inputs = [
        ("Bars\n64 × 16", "#E8EFF7", "#2F5D8C"),
        ("L1\n240 × 16", "#E7F4F1", "#0F766E"),
        ("L2\n240 × 2 × 10 × 12", "#F4EEE5", "#9A6724"),
        ("Quality\n16", "#F1EDF7", "#70558C"),
    ]
    encoders = [
        "4-frame patches\n3 causal blocks",
        "4-frame patches\n3 causal blocks",
        "side/rank hierarchy\n→ patches → 3 blocks",
        "quality conditioner",
    ]
    for y, (label, face, edge), encoder in zip(input_y, inputs, encoders):
        rect(ax, (0.05, y), 1.85, 0.78, label, face=face, edge=edge, weight="bold")
        rect(ax, (2.35, y), 2.15, 0.78, encoder, face="#F8FAFC", edge=edge, fontsize=9.5)
        arrow(ax, (1.9, y + 0.39), (2.35, y + 0.39), color=edge)

    rect(ax, (5.15, 2.6), 2.05, 2.6, "Fused causal memory\n184 tokens × 128\n\nquality-aware\nmodality gating", face="#EEF3F7", edge="#607080", weight="bold")
    for y in input_y:
        arrow(ax, (4.5, y + 0.39), (5.15, 3.9), color="#607080")

    rect(ax, (7.85, 3.15), 1.8, 1.5, "5 horizon queries\n2 decoder blocks\nwidth 128", face="#E8EFF7", edge="#2F5D8C", weight="bold")
    arrow(ax, (7.2, 3.9), (7.85, 3.9), color="#2F5D8C")
    rect(ax, (10.15, 3.15), 1.75, 1.5, "Output [5,17]\n15 quantiles\n2 presence logits", face="#F9ECE9", edge="#B44C3B", weight="bold")
    arrow(ax, (9.65, 3.9), (10.15, 3.9), color="#B44C3B")

    rect(ax, (7.1, 0.75), 4.8, 1.45, "Five targets per horizon\ntotal POC • buy POC • sell POC • log-volume residual • delta fraction\nq10 / q50 / q90", face="#F8FAFC", edge="#607080", fontsize=9.8)
    arrow(ax, (11.0, 3.15), (10.2, 2.2), color="#607080")
    ax.text(0.05, 0.45, "Trainable parameters: 2,089,554", fontsize=12, weight="bold", color="#102A43")
    ax.text(0.05, 0.12, "Objective: weighted pinball + median Huber + presence BCE + scheduled auxiliary loss", fontsize=10.5, color="#5F6B78")
    save_figure(fig, path)


def figure_score_heatmap(path: Path) -> None:
    values = np.array(
        [
            [6.4068, 6.1584, 5.3419, 1.9157, -5.7273],
            [5.9117, 6.3689, 5.1958, -3.7062, -10.6858],
            [4.9555, 5.7222, 4.7519, -9.2200, -16.8138],
            [3.8155, 5.1001, 4.2883, -15.1326, -23.2551],
            [2.6231, 4.8357, 3.5302, -22.4325, -29.5699],
        ]
    )
    fig, ax = plt.subplots(figsize=(10.7, 5.7))
    image = ax.imshow(values, cmap="RdYlGn", vmin=-30, vmax=30, aspect="auto")
    ax.set_title("Pinball-loss improvement by target and horizon", loc="left", pad=18, weight="bold")
    ax.text(-0.5, -0.95, "Relative to empirical marginal baseline; positive values favor CHMT-Q", color="#5F6B78", transform=ax.transData)
    ax.set_xticks(range(5), ["Total POC", "Buy POC", "Sell POC", "Log volume", "Delta"])
    ax.set_yticks(range(5), ["H1", "H2", "H3", "H4", "H5"])
    ax.set_xlabel("Forecast target")
    ax.set_ylabel("One-minute horizon")
    for row in range(5):
        for col in range(5):
            color = "white" if abs(values[row, col]) > 15 else "#172033"
            ax.text(col, row, f"{values[row, col]:+.1f}%", ha="center", va="center", color=color, weight="bold", fontsize=10)
    cbar = fig.colorbar(image, ax=ax, fraction=0.035, pad=0.035)
    cbar.set_label("Relative improvement (%)")
    ax.set_xticks(np.arange(-0.5, 5, 1), minor=True)
    ax.set_yticks(np.arange(-0.5, 5, 1), minor=True)
    ax.grid(which="minor", color="white", linewidth=2)
    ax.tick_params(which="minor", bottom=False, left=False)
    save_figure(fig, path)


def figure_h5_diagnostics(path: Path) -> None:
    labels = ["Pearson", "Spearman", "Sign MCC", "Median\ndispersion", "80% interval\ncoverage", "Intervals\ncrossing zero"]
    values = [0.07745, 0.07533, 0.04809, 0.15130, 0.81150, 1.0]
    colors = ["#2F5D8C", "#2F5D8C", "#C88918", "#C88918", "#0F766E", "#B44C3B"]
    fig, ax = plt.subplots(figsize=(11, 5.35))
    bars = ax.bar(range(len(values)), values, color=colors, width=0.66)
    ax.axhline(0.8, color="#607080", linestyle="--", linewidth=1.0)
    ax.text(5.45, 0.815, "nominal 80%", ha="right", va="bottom", color="#5F6B78", fontsize=9)
    ax.set_ylim(0, 1.12)
    ax.set_xticks(range(len(labels)), labels)
    ax.set_ylabel("Metric value / fraction")
    ax.set_title("Horizon-five total-POC signal: association is weak, intervals are nonselective", loc="left", pad=18, weight="bold")
    ax.grid(axis="y", color="#E2E8F0", linewidth=0.8)
    ax.spines[["top", "right"]].set_visible(False)
    for bar, value in zip(bars, values):
        ax.text(bar.get_x() + bar.get_width() / 2, value + 0.025, f"{value:.3f}", ha="center", va="bottom", weight="bold")
    ax.text(0.01, -0.23, "Median dispersion = predicted standard deviation / realized standard deviation.", transform=ax.transAxes, color="#5F6B78")
    save_figure(fig, path)


def figure_cost_feasibility(path: Path) -> None:
    fig, ax = plt.subplots(figsize=(11.2, 5.5))
    categories = ["Long-flat\ngross edge", "Long-short\ngross edge", "Long-flat\nbreak-even / side", "Optimistic\ncost / side", "Base\ncost / side", "Stress\ncost / side"]
    values = [0.4453, 0.4301, 0.2227, 9.0, 12.0, 20.0]
    colors = ["#0F766E", "#2F5D8C", "#607080", "#C88918", "#B44C3B", "#7A2E23"]
    bars = ax.bar(range(len(values)), values, color=colors, width=0.66)
    ax.set_yscale("log")
    ax.set_ylim(0.15, 30)
    ax.set_xticks(range(len(categories)), categories)
    ax.set_ylabel("Basis points (log scale)")
    ax.set_title("Measured gross edge is orders of magnitude below tested costs", loc="left", pad=18, weight="bold")
    ax.grid(axis="y", color="#E2E8F0", which="both", linewidth=0.8)
    ax.spines[["top", "right"]].set_visible(False)
    for bar, value in zip(bars, values):
        ax.text(bar.get_x() + bar.get_width() / 2, value * 1.12, f"{value:g} bps", ha="center", va="bottom", weight="bold", fontsize=9.5)
    ax.annotate(
        "Base one-way cost is 53.89×\nthe long-flat break-even cost",
        xy=(4, 12),
        xytext=(2.45, 22),
        arrowprops={"arrowstyle": "-|>", "color": "#B44C3B"},
        color="#7A2E23",
        weight="bold",
        ha="center",
    )
    save_figure(fig, path)


def figure_failure_map(path: Path) -> None:
    labels = [
        "Target is not trading return",
        "Gross edge far below costs",
        "All intervals cross zero",
        "Weak, shrunk H5 medians",
        "Dependent overlapping windows",
        "One seed / fold; four forward days",
        "Degenerate presence labels",
        "Saturated delta quantiles",
        "Persistent gradient clipping",
        "Sparse-feature scale clipping",
        "Weak registered baselines",
        "Simplified execution proxy",
    ]
    scores = np.arange(12, 0, -1)
    status = ["Decisive"] * 2 + ["Major"] * 2 + ["Limitation"] * 3 + ["Pathology"] + ["Needs ablation"] * 2 + ["Limitation"] * 2
    color_map = {
        "Decisive": "#B44C3B",
        "Major": "#C88918",
        "Limitation": "#607080",
        "Pathology": "#9A6724",
        "Needs ablation": "#2F5D8C",
    }
    colors = [color_map[item] for item in status]
    fig, ax = plt.subplots(figsize=(11.3, 7.0))
    y = np.arange(len(labels))
    bars = ax.barh(y, scores, color=colors, height=0.66)
    for index in [8, 9]:
        bars[index].set_hatch("///")
        bars[index].set_edgecolor("white")
    ax.set_yticks(y, labels)
    ax.invert_yaxis()
    ax.set_xlim(0, 13.8)
    ax.set_xticks([])
    ax.set_title("Failure-factor audit: observed causes are separated from ablation hypotheses", loc="left", pad=18, weight="bold")
    ax.spines[:].set_visible(False)
    for bar, label in zip(bars, status):
        ax.text(bar.get_width() + 0.15, bar.get_y() + bar.get_height() / 2, label, va="center", fontsize=9.5, color="#465362", weight="bold")
    ax.text(0, 12.5, "Rank encodes audit priority, not a causal effect size. Hatched factors require controlled ablation.", color="#5F6B78")
    save_figure(fig, path)


def figure_next_study(path: Path) -> None:
    fig, ax = plt.subplots(figsize=(12, 5.8))
    ax.set_xlim(0, 12)
    ax.set_ylim(0, 6)
    ax.axis("off")
    ax.text(0, 5.72, "Corrective protocol: every promotion step has a falsifiable gate", fontsize=17, weight="bold")
    stages = [
        (0.05, "1  Data integrity", "≥240 sessions\nsequence + receive time\nlicense review", "#E8EFF7", "#2F5D8C"),
        (2.45, "2  Forecast baselines", "direct returns\nstrong simple models\n≥5 seeds / 5 folds", "#E7F4F1", "#0F766E"),
        (4.85, "3  Cost feasibility", "executable quotes\nturnover + impact\nsafety margin", "#F4EEE5", "#9A6724"),
        (7.25, "4  DQN eligibility", "Double + dueling\nnegative controls\n≥10 seeds", "#F1EDF7", "#70558C"),
        (9.65, "5  Confirmation", "new ≥60-session holdout\nthen ≥60 shadow days\nno capital before gates", "#F9ECE9", "#B44C3B"),
    ]
    for x, title, body, face, edge in stages:
        rect(ax, (x, 2.35), 2.05, 2.15, f"{title}\n\n{body}", face=face, edge=edge, weight="bold", fontsize=9.5)
    for index in range(4):
        arrow(ax, (stages[index][0] + 2.05, 3.42), (stages[index + 1][0], 3.42), color="#607080")
        ax.text((stages[index][0] + 2.05 + stages[index + 1][0]) / 2, 4.62, "PASS", ha="center", fontsize=8.5, color="#0F766E", weight="bold")
    ax.text(0.05, 1.55, "STOP CONDITIONS", fontsize=10, weight="bold", color="#7A2E23")
    ax.text(0.05, 1.12, "look-ahead • unsupported actions • negative-control profit • Q-value divergence • edge below costs • unstable seeds", fontsize=10.5, color="#7A2E23")
    ax.text(0.05, 0.45, "The current experiment is before Gate 2: direct-return utility and strong-baseline superiority are not established.", fontsize=11.5, weight="bold", color="#102A43")
    save_figure(fig, path)


def build_document(source: Path) -> Document:
    document = Document()
    configure_document(document)
    renderer = ManuscriptRenderer(document, source.parent)
    renderer.render(source.read_text(encoding="utf-8"))
    add_core_properties(document)
    return document


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.78)
    section.left_margin = Inches(0.92)
    section.right_margin = Inches(0.92)
    section.header_distance = Inches(0.34)
    section.footer_distance = Inches(0.35)
    section.different_first_page_header_footer = True

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.13
    normal.paragraph_format.widow_control = True

    title = styles["Title"]
    title.font.name = "Times New Roman"
    title.font.size = Pt(25)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(NAVY)
    title.paragraph_format.space_before = Pt(62)
    title.paragraph_format.space_after = Pt(14)
    title.paragraph_format.keep_with_next = True

    for name, size, color, before, after in [
        ("Heading 1", 15, NAVY, 16, 6),
        ("Heading 2", 12, TEAL, 12, 4),
        ("Heading 3", 10.5, INK, 9, 3),
    ]:
        style = styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    subtitle = add_style(document, "Manuscript Subtitle", WD_STYLE_TYPE.PARAGRAPH)
    subtitle.font.name = "Times New Roman"
    subtitle.font.size = Pt(15)
    subtitle.font.italic = True
    subtitle.font.color.rgb = RGBColor.from_string(TEAL)
    subtitle.paragraph_format.space_after = Pt(22)
    subtitle.paragraph_format.keep_with_next = True

    abstract = add_style(document, "Abstract Body", WD_STYLE_TYPE.PARAGRAPH)
    abstract.font.name = "Times New Roman"
    abstract.font.size = Pt(9.8)
    abstract.paragraph_format.left_indent = Inches(0.35)
    abstract.paragraph_format.right_indent = Inches(0.35)
    abstract.paragraph_format.space_after = Pt(6)
    abstract.paragraph_format.line_spacing = 1.08

    caption = styles["Caption"]
    caption.font.name = "Times New Roman"
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(MUTED)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.keep_together = True

    table_caption = add_style(document, "Table Caption", WD_STYLE_TYPE.PARAGRAPH)
    table_caption.font.name = "Times New Roman"
    table_caption.font.size = Pt(9)
    table_caption.font.bold = True
    table_caption.font.color.rgb = RGBColor.from_string(NAVY)
    table_caption.paragraph_format.space_before = Pt(7)
    table_caption.paragraph_format.space_after = Pt(3)
    table_caption.paragraph_format.keep_with_next = True

    equation = add_style(document, "Equation Text", WD_STYLE_TYPE.PARAGRAPH)
    equation.font.name = "Cambria Math"
    equation.font.size = Pt(10.5)
    equation.paragraph_format.space_before = Pt(5)
    equation.paragraph_format.space_after = Pt(5)
    equation.paragraph_format.keep_together = True

    reference = add_style(document, "Reference Entry", WD_STYLE_TYPE.PARAGRAPH)
    reference.font.name = "Times New Roman"
    reference.font.size = Pt(8.8)
    reference.paragraph_format.left_indent = Inches(0.28)
    reference.paragraph_format.first_line_indent = Inches(-0.28)
    reference.paragraph_format.space_after = Pt(2.2)
    reference.paragraph_format.line_spacing = 1.03

    bullet = styles["List Bullet"]
    bullet.font.name = "Times New Roman"
    bullet.font.size = Pt(10.5)
    bullet.paragraph_format.left_indent = Inches(0.28)
    bullet.paragraph_format.first_line_indent = Inches(-0.17)
    bullet.paragraph_format.space_after = Pt(2.5)

    numbered = styles["List Number"]
    numbered.font.name = "Times New Roman"
    numbered.font.size = Pt(10.5)
    numbered.paragraph_format.left_indent = Inches(0.32)
    numbered.paragraph_format.first_line_indent = Inches(-0.20)
    numbered.paragraph_format.space_after = Pt(3)

    configure_header_footer(section)


def add_style(document: Document, name: str, style_type: WD_STYLE_TYPE):
    try:
        return document.styles[name]
    except KeyError:
        return document.styles.add_style(name, style_type)


def configure_header_footer(section) -> None:
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.tab_stops.add_tab_stop(
        Inches(6.62), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.SPACES
    )
    left = paragraph.add_run("CHMT-Q NEGATIVE-RESULTS STUDY")
    paragraph.add_run("\t")
    right = paragraph.add_run("ANONYMOUS MANUSCRIPT")
    for run in [left, right]:
        run.font.name = "Arial"
        run.font.size = Pt(7)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(MUTED)
    properties = paragraph._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "5")
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), GRID)
    border.append(bottom)
    properties.append(border)

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(2)
    run = paragraph.add_run("Development research • Not for capital allocation    |    ")
    run.font.name = "Arial"
    run.font.size = Pt(7.5)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    add_field(paragraph, "PAGE")


def add_core_properties(document: Document) -> None:
    properties = document.core_properties
    properties.title = "When Better Proper Scores Do Not Imply Tradable Edge"
    properties.subject = "Leakage-aware failure analysis of CHMT-Q"
    properties.author = "Anonymous"
    properties.last_modified_by = "Anonymous"
    properties.keywords = "limit order book, quantile forecasting, Transformer, negative results"
    properties.comments = "Venue-neutral anonymous manuscript; quantitative claims require author verification."


class ManuscriptRenderer:
    def __init__(self, document: Document, base_dir: Path) -> None:
        self.document = document
        self.base_dir = base_dir
        self.figure_number = 0
        self.table_number = 0
        self.equation_number = 0
        self.section_title = ""
        self.cover_subtitle_pending = True
        self.abstract_mode = False
        self.appendix_started = False

    def render(self, markdown: str) -> None:
        lines = markdown.splitlines()
        index = 0
        while index < len(lines):
            raw = lines[index]
            stripped = raw.strip()
            if not stripped:
                index += 1
                continue
            if stripped == "{PAGEBREAK}":
                self.document.add_page_break()
                index += 1
                continue
            if stripped.startswith("{FIGURE:"):
                self.add_figure(stripped)
                index += 1
                continue
            if stripped.startswith("{EQUATION:"):
                self.add_equation(stripped)
                index += 1
                continue
            if stripped.startswith("|") and stripped.endswith("|"):
                block = []
                while index < len(lines) and lines[index].strip().startswith("|"):
                    block.append(lines[index].strip())
                    index += 1
                self.add_table(block)
                continue
            heading = re.match(r"^(#{1,4})\s+(.+)$", stripped)
            if heading:
                self.add_heading(len(heading.group(1)), heading.group(2))
                index += 1
                continue
            if stripped.startswith("- "):
                while index < len(lines) and lines[index].strip().startswith("- "):
                    paragraph = self.document.add_paragraph(style="List Bullet")
                    add_inline(paragraph, lines[index].strip()[2:])
                    paragraph.paragraph_format.keep_together = True
                    index += 1
                continue
            if re.match(r"^\d+\.\s+", stripped):
                reference_mode = self.section_title == "References"
                while index < len(lines) and re.match(r"^\d+\.\s+", lines[index].strip()):
                    text = lines[index].strip()
                    if reference_mode:
                        paragraph = self.document.add_paragraph(style="Reference Entry")
                        add_inline(paragraph, text)
                    else:
                        paragraph = self.document.add_paragraph(style="List Number")
                        add_inline(paragraph, re.sub(r"^\d+\.\s+", "", text))
                    paragraph.paragraph_format.keep_together = True
                    index += 1
                continue

            paragraph_lines = [stripped]
            index += 1
            while index < len(lines):
                candidate = lines[index].strip()
                if not candidate or is_block_start(candidate):
                    break
                paragraph_lines.append(candidate)
                index += 1
            text = " ".join(paragraph_lines)
            style = "Abstract Body" if self.abstract_mode and not text.startswith("**Keywords:") else "Normal"
            paragraph = self.document.add_paragraph(style=style)
            add_inline(paragraph, text)
            paragraph.paragraph_format.keep_together = len(text) < 420

    def add_heading(self, level: int, text: str) -> None:
        clean = strip_markdown(text)
        if level == 1:
            paragraph = self.document.add_paragraph(style="Title")
            add_inline(paragraph, text)
            self.cover_subtitle_pending = True
            return
        if self.cover_subtitle_pending and level == 2:
            paragraph = self.document.add_paragraph(style="Manuscript Subtitle")
            add_inline(paragraph, text)
            self.cover_subtitle_pending = False
            return

        self.cover_subtitle_pending = False
        if clean == "References":
            self.document.add_page_break()
        if clean.startswith("Appendix A") and not self.appendix_started:
            self.document.add_page_break()
            self.appendix_started = True
        style_level = min(max(level - 1, 1), 3)
        paragraph = self.document.add_paragraph(style=f"Heading {style_level}")
        add_inline(paragraph, text)
        self.section_title = clean
        self.abstract_mode = clean == "Abstract"
        if clean != "Abstract" and style_level == 1:
            self.abstract_mode = False

    def add_figure(self, marker: str) -> None:
        match = re.match(r"^\{FIGURE:([^|]+)\|([^|]+)\|(.+)\}$", marker)
        if not match:
            raise ValueError(f"Invalid figure marker: {marker}")
        filename, caption, alt_text = match.groups()
        path = self.base_dir / "assets" / filename
        if not path.is_file():
            raise FileNotFoundError(path)
        self.figure_number += 1
        paragraph = self.document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(6)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.keep_together = True
        run = paragraph.add_run()
        inline = run.add_picture(str(path), width=Inches(6.35))
        doc_pr = inline._inline.docPr
        doc_pr.set("descr", alt_text)
        doc_pr.set("title", f"Figure {self.figure_number}")
        caption_paragraph = self.document.add_paragraph(style="Caption")
        add_inline(caption_paragraph, caption)
        caption_paragraph.paragraph_format.keep_together = True

    def add_equation(self, marker: str) -> None:
        equation = marker[len("{EQUATION:") : -1]
        self.equation_number += 1
        paragraph = self.document.add_paragraph(style="Equation Text")
        paragraph.paragraph_format.tab_stops.add_tab_stop(
            Inches(3.30), WD_TAB_ALIGNMENT.CENTER, WD_TAB_LEADER.SPACES
        )
        paragraph.paragraph_format.tab_stops.add_tab_stop(
            Inches(6.40), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.SPACES
        )
        paragraph.add_run("\t")
        run = paragraph.add_run(equation)
        run.font.name = "Cambria Math"
        run.font.size = Pt(9.8)
        paragraph.add_run("\t")
        number = paragraph.add_run(f"({self.equation_number})")
        number.font.name = "Times New Roman"
        number.font.size = Pt(9.5)

    def add_table(self, block: list[str]) -> None:
        rows = [parse_table_row(line) for line in block]
        if len(rows) < 2 or not all(re.fullmatch(r":?-{3,}:?", item.replace(" ", "")) for item in rows[1]):
            raise ValueError(f"Malformed Markdown table near: {block[0]}")
        data = [rows[0]] + rows[2:]
        columns = len(data[0])
        if any(len(row) != columns for row in data):
            raise ValueError(f"Inconsistent table width near: {block[0]}")

        self.table_number += 1
        caption_text = TABLE_CAPTIONS[self.table_number - 1] if self.table_number <= len(TABLE_CAPTIONS) else "Supplementary results."
        caption = self.document.add_paragraph(style="Table Caption")
        caption.add_run(f"Table {self.table_number}. {caption_text}")

        table = self.document.add_table(rows=len(data), cols=columns)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        widths = table_widths(self.table_number, columns)
        for col_index, width in enumerate(widths):
            table.columns[col_index].width = Inches(width)
        for row_index, values in enumerate(data):
            row = table.rows[row_index]
            set_row_cant_split(row)
            if row_index == 0:
                set_repeat_table_header(row)
            for col_index, value in enumerate(values):
                cell = row.cells[col_index]
                cell.width = Inches(widths[col_index])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_cell_margins(cell, top=55, bottom=55, left=60, right=60)
                paragraph = cell.paragraphs[0]
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                if row_index == 0:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                elif looks_numeric(value):
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                add_inline(paragraph, value, table_cell=True)
                for run in paragraph.runs:
                    run.font.size = Pt(7.5 if self.table_number == 8 else 8.2)
                    if row_index == 0:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor.from_string(WHITE)
                if row_index == 0:
                    set_cell_shading(cell, NAVY)
                elif row_index % 2 == 0:
                    set_cell_shading(cell, PALE)
        set_table_borders(
            table,
            {
                "top": {"val": "single", "sz": "6", "color": NAVY},
                "bottom": {"val": "single", "sz": "6", "color": NAVY},
                "insideH": {"val": "single", "sz": "3", "color": GRID},
                "insideV": {"val": "single", "sz": "3", "color": GRID},
            },
        )
        spacer = self.document.add_paragraph()
        spacer.paragraph_format.space_after = Pt(1)
        spacer.paragraph_format.space_before = Pt(0)


def is_block_start(text: str) -> bool:
    return bool(
        text == "{PAGEBREAK}"
        or text.startswith("{FIGURE:")
        or text.startswith("{EQUATION:")
        or text.startswith("#")
        or text.startswith("|")
        or text.startswith("- ")
        or re.match(r"^\d+\.\s+", text)
    )


def strip_markdown(text: str) -> str:
    return re.sub(r"[*`]", "", text)


def add_inline(paragraph, text: str, *, table_cell: bool = False) -> None:
    pattern = re.compile(r"(\*\*.+?\*\*|`.+?`|\*[^*]+?\*|https?://[^\s]+)")
    cursor = 0
    for match in pattern.finditer(text):
        if match.start() > cursor:
            add_text_run(paragraph, text[cursor : match.start()], table_cell=table_cell)
        token = match.group(0)
        if token.startswith("**"):
            run = add_text_run(paragraph, token[2:-2], table_cell=table_cell)
            run.bold = True
        elif token.startswith("`"):
            value = token[1:-1]
            if table_cell and len(value) > 24:
                value = "\u200b".join(value[index : index + 8] for index in range(0, len(value), 8))
            run = add_text_run(paragraph, value, table_cell=table_cell)
            run.font.name = "Courier New"
            run.font.size = Pt(7.1 if table_cell else 9)
            run.font.color.rgb = RGBColor.from_string(NAVY)
        elif token.startswith("*"):
            run = add_text_run(paragraph, token[1:-1], table_cell=table_cell)
            run.italic = True
        else:
            trailing = ""
            url = token
            while url and url[-1] in ".,;:)":
                trailing = url[-1] + trailing
                url = url[:-1]
            add_hyperlink(paragraph, url, url)
            if trailing:
                add_text_run(paragraph, trailing, table_cell=table_cell)
        cursor = match.end()
    if cursor < len(text):
        add_text_run(paragraph, text[cursor:], table_cell=table_cell)


def add_text_run(paragraph, text: str, *, table_cell: bool = False):
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(8.2 if table_cell else 10.5)
    run.font.color.rgb = RGBColor.from_string(INK)
    return run


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    relationship_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    properties.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(underline)
    run.append(properties)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def parse_table_row(line: str) -> list[str]:
    return [item.strip() for item in line.strip().strip("|").split("|")]


def looks_numeric(value: str) -> bool:
    plain = strip_markdown(value).strip()
    return bool(re.match(r"^[+−-]?[\d≥].*[%×]?$", plain))


def table_widths(table_number: int, columns: int) -> list[float]:
    overrides = {
        1: [1.25, 1.6, 1.75, 2.0],
        2: [1.75, 2.15, 2.7],
        3: [2.7, 3.9],
        4: [2.55, 1.05, 1.05, 1.95],
        5: [1.8, 0.7, 1.45, 1.45, 1.2],
        6: [0.45, 1.6, 1.45, 3.1],
        7: [1.35, 3.35, 1.9],
        8: [1.35, 3.65, 1.6],
        9: [1.65, 3.2, 1.75],
        10: [2.05, 1.55, 3.0],
    }
    if table_number in overrides:
        return overrides[table_number]
    return [6.6 / columns] * columns


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, *, top: int, bottom: int, left: int, right: int) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for key, value in {"top": top, "bottom": bottom, "start": left, "end": right}.items():
        node = margins.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def set_row_cant_split(row) -> None:
    properties = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    properties.append(cant_split)


def set_table_borders(table, borders: dict[str, dict[str, str]]) -> None:
    properties = table._tbl.tblPr
    element = properties.first_child_found_in("w:tblBorders")
    if element is None:
        element = OxmlElement("w:tblBorders")
        properties.append(element)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        node = element.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            element.append(node)
        definition = borders.get(edge, {"val": "nil"})
        for key, value in definition.items():
            node.set(qn(f"w:{key}"), value)


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = f" {instruction} "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, text, separate, end])


if __name__ == "__main__":
    raise SystemExit(main())
